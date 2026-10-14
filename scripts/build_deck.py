"""Generate a draft PowerPoint deck (docs/MCPShield.pptx) from the docs and eval figures.

Plain professional styling, no emojis. Figures are pulled from docs/figures/ when present (PNG only; .mmd
sources and .html interactive charts are skipped). Run after `make eval` and `make figures` for charts.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from mcpshield.config import get_settings

TITLE = "MCPShield"
SUBTITLE = "A multi-agent runtime firewall for MCP-connected AI agents"


def _bullets(slide, title: str, points: list[str]) -> None:
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, point in enumerate(points):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.text = point
        para.font.size = Pt(18)


def _picture_slide(prs: Presentation, title: str, image: Path, caption: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.add_picture(str(image), Inches(0.7), Inches(1.6), width=Inches(8.6))
    if caption:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(6.9), Inches(8.6), Inches(0.5))
        box.text_frame.text = caption


def build() -> Path:
    settings = get_settings()
    figs = settings.paths.figures_dir
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = TITLE
    title_slide.placeholders[1].text = SUBTITLE

    content = prs.slide_layouts[1]
    _bullets(
        prs.slides.add_slide(content),
        "Problem",
        [
            "Enterprises connect AI agents to tools via MCP servers.",
            "The agent trusts tool metadata and tool outputs as context.",
            "A malicious/compromised MCP server can hijack the agent:",
            "  poisoning, shadowing, rug pulls, indirect injection, split payloads, exfiltration.",
        ],
    )
    _bullets(
        prs.slides.add_slide(content),
        "Existing challenge",
        [
            "Today: manual tool review and static allowlists.",
            "Misses hidden instructions and split payloads.",
            "Cannot see rug pulls (trusted tool changes later).",
            "Misses runtime attacks (injection in outputs, tainted arguments).",
            "No audit trail for compliance.",
        ],
    )
    _bullets(
        prs.slides.add_slide(content),
        "Proposed solution",
        [
            "A transparent proxy between the host agent and MCP servers.",
            "Every MCP event flows through a LangGraph multi-agent pipeline.",
            "Decides ALLOW / SANITIZE / BLOCK / ESCALATE with an explanation.",
            "Hash-pins tools, scans with RAG over a threat KB, enforces policy, tracks taint.",
        ],
    )
    _bullets(
        prs.slides.add_slide(content),
        "Agent roles",
        [
            "Orchestrator - Planner, utility-based (fast vs deep path).",
            "Registry Auditor - Research, model-based (hash pinning, rug-pull, RAG).",
            "Intent Verifier - goal-based (policy, sensitive data, taint).",
            "Output Sanitizer - reflex + LLM (injection in outputs).",
            "Critic - utility-based (risk aggregation, self-consistency, guardrail).",
            "Auditor - monitoring (audit log, explainability).",
        ],
    )
    _bullets(
        prs.slides.add_slide(content),
        "Orchestration workflow",
        [
            "Perception: proxy turns MCP traffic into events.",
            "Reasoning: plan -> rule detectors -> (grey-zone) LLM judge -> Critic.",
            "Fast path for clear cases; deep path only in the grey zone (latency/risk trade-off).",
            "Action: proxy applies the decision; Auditor records it.",
        ],
    )

    for name, title, caption in [
        ("architecture_system", "Architecture", "System diagram (see docs/architecture.md for the Mermaid source)."),
        ("asr_by_config", "Results: attack success rate", "ASR drops from the shield-OFF baseline to near zero."),
        ("asr_per_class", "Results: per-class ASR", "Per attack class, across ablations."),
        ("fpr_by_config", "Results: false positives", "FPR on benign tasks stays low."),
        ("latency", "Results: added latency", "Deep path adds latency only in the grey zone."),
    ]:
        png = figs / f"{name}.png"
        if png.exists():
            _picture_slide(prs, title, png, caption)
        else:
            _bullets(
                prs.slides.add_slide(content),
                title,
                [caption, f"(figure {name}.png not found; run make eval && make figures)"],
            )

    _bullets(
        prs.slides.add_slide(content),
        "Validation",
        [
            "Deterministic mock backend: reproducible, offline, worst-case obedient agent.",
            "Shield-OFF ASR is 100% by construction; the result is the drop with the shield on.",
            "Ground truth from side effects (ledger, outbox, exfil sink); results are order-independent.",
            "Every decision has a human-readable explanation in the audit log and UI.",
        ],
    )
    _bullets(
        prs.slides.add_slide(content),
        "Limitations and future scope",
        [
            "Limitations: LLM judge is injectable; small models misclassify; attacks are synthetic;",
            "  deep path adds latency; taint tracking is coarse; offline embeddings are weaker.",
            "Future: full taint tracking, signed manifests, fine-tuned classifier, A2A protection,",
            "  adaptive red-teaming, sidecar deployment with monitoring.",
        ],
    )
    _bullets(
        prs.slides.add_slide(content),
        "Conclusion",
        [
            "MCPShield lets enterprises adopt untrusted MCP tools safely.",
            "It blocks runtime attacks that static review misses, at a low false-positive cost.",
            "It leaves a full, explainable audit trail for compliance.",
        ],
    )

    out = settings.paths.data_dir.parent / "docs" / "MCPShield.pptx"
    prs.save(str(out))
    return out


def main() -> None:
    out = build()
    print(f"deck written to {out}")


if __name__ == "__main__":
    main()
